"""Excel export laid out like the old sheet (PLAN.md §11.3): one sheet per budget, with
Person 1 (Monzo then Starling), Person 2 RBS and Joint Billing side by side, then the transfers
and surpluses."""

import re
from datetime import date
from decimal import Decimal
from fractions import Fraction
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.models import BudgetVersion
from app.domain.versions import ResolvedBudget, calculate, resolve
from app.engine.money import round_pence

MONEY = "£#,##0.00;[Red]-£#,##0.00"
BAND = PatternFill("solid", fgColor="1F4A73")
WHITE_BOLD = Font(bold=True, color="F1F4F7")
BOLD = Font(bold=True)
RULE = Border(top=Side(style="thin", color="8D958E"))


def _pounds(pence: Fraction | int) -> Decimal:
    return Decimal(round_pence(pence)) / 100


def _title(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "", name)[:31] or "Budget"
    title, n = base, 2
    while title in used:
        title = f"{base[:28]}({n})"
        n += 1
    used.add(title)
    return title


def _block(
    ws: Worksheet,
    col: int,
    row: int,
    heading: str,
    resolved: ResolvedBudget,
    accounts: list[str],
    amounts: dict[str, int],
) -> int:
    """One account column: Day | Item | Out. Returns the next free row."""
    ws.cell(row, col, heading).font = BOLD
    for i, label in enumerate(("Day", "Item", "Out")):
        ws.cell(row + 1, col + i, label).font = Font(bold=True, color="4A535B")
    r = row + 2
    total = Fraction(0)
    for account in accounts:
        if len(accounts) > 1:
            ws.cell(
                r, col + 1, next(a.name for a in resolved.accounts if a.key == account)
            ).font = BOLD
            r += 1
        lines = sorted(
            (ln for ln in resolved.lines if ln.account == account),
            key=lambda ln: (ln.day_of_month or 99, ln.name),
        )
        for ln in lines:
            value: Fraction | int = (
                amounts.get(ln.key, 0) if ln.kind == "transfer" else ln.monthly()
            )
            ws.cell(r, col, ln.day_of_month)
            ws.cell(r, col + 1, ln.label)
            cell = ws.cell(r, col + 2, _pounds(value))
            cell.number_format = MONEY
            total += value
            r += 1
    ws.cell(r, col + 1, "Total out").font = BOLD
    cell = ws.cell(r, col + 2, _pounds(total))
    cell.number_format, cell.font = MONEY, BOLD
    for c in range(col, col + 3):
        ws.cell(r, c).border = RULE
    return r + 1


def _sheet(ws: Worksheet, resolved: ResolvedBudget) -> None:
    result = calculate(resolved)
    amounts = {t.line_key: t.effective_pence for t in result.transfers}
    owners = {a.key: a for a in resolved.accounts}

    ws.cell(1, 1, f"A Wee Bawbee: {resolved.version_name}").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"As of {resolved.as_of:%d %B %Y}. Monthly amounts; annual bills shown ÷ 12.")

    def keys(owner: str | None, roles: tuple[str, ...]) -> list[str]:
        return [a.key for a in resolved.accounts if a.owner == owner and a.role in roles]

    columns = [
        ("Person 1", keys("person_1", ("salary_hub",)) + keys("person_1", ("subscriptions",))),
        ("Person 2", keys("person_2", ("salary_hub",))),
        ("Joint Billing", [a.key for a in resolved.accounts if a.role == "joint_bills"]),
    ]
    bottom = 4
    for i, (heading, accounts) in enumerate(columns):
        if accounts:
            bottom = max(bottom, _block(ws, 1 + i * 4, 4, heading, resolved, accounts, amounts))

    r = bottom + 2
    ws.cell(r, 1, "Transfers").font = BOLD
    for i, label in enumerate(("From", "To", "Suggested", "Actual")):
        ws.cell(r + 1, 1 + i, label).font = Font(bold=True, color="4A535B")
    r += 2
    for t in result.transfers:
        ws.cell(r, 1, owners[t.from_account].name)
        ws.cell(r, 2, owners[t.to_account].name)
        ws.cell(r, 3, _pounds(t.suggested_pence)).number_format = MONEY
        if t.actual_pence is not None:
            ws.cell(r, 4, _pounds(t.actual_pence)).number_format = MONEY
        r += 1

    r += 1
    ws.cell(r, 1, "Surplus").font = BOLD
    r += 1
    for person, value in result.personal_surplus.items():
        ws.cell(r, 1, person.replace("_", " ").title())
        ws.cell(r, 3, _pounds(value)).number_format = MONEY
        r += 1
    for label, value in (("Joint", result.joint_surplus), ("Household", result.household_surplus)):
        ws.cell(r, 1, label).font = BOLD if label == "Household" else Font()
        ws.cell(r, 3, _pounds(value)).number_format = MONEY
        r += 1

    for col, width in zip("ABCDEFGHIJK", (6, 34, 12, 3, 6, 30, 12, 3, 6, 40, 12), strict=True):
        ws.column_dimensions[col].width = width
    for c in (1, 5, 9):
        header = ws.cell(4, c)
        for cc in range(c, c + 3):
            ws.cell(4, cc).fill = BAND
        header.font = WHITE_BOLD
        header.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A6"


def workbook(session: Session, as_of: date) -> bytes:
    versions = session.scalars(
        select(BudgetVersion)
        .where(BudgetVersion.status.in_(("active", "draft")))
        .order_by(BudgetVersion.type != "live", BudgetVersion.id)
    ).all()
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    used: set[str] = set()
    for version in versions:
        _sheet(wb.create_sheet(_title(version.name, used)), resolve(session, version, as_of))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
