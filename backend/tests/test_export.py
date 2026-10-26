from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.importers.seed import ImportResult


def test_json_backup_has_the_data_but_no_login_secrets(
    client: TestClient, imported: ImportResult
) -> None:
    r = client.get("/api/export/backup.json")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    body = r.json()
    assert body["schema_revision"] == "0007"
    tables = body["tables"]
    assert {
        "budget_version",
        "budget_line",
        "budget_line_version",
        "override",
        "audit_log",
        "bank_transaction",
        "import_batch",
        "category_rule",
        "merchant_alias",
        "inbox_item",
        "bank_connection",
        "bank_account_link",
        "sync_run",
        "recurring_candidate",
        "notification",
        "ha_token",
    } <= set(tables)
    assert len(tables["budget_version"]) == 2
    assert "auth_session" not in tables
    assert "login_attempt" not in tables
    assert "bank_link_flow" not in tables
    assert all("password_hash" not in u for u in tables["user"])
    assert {u["username"] for u in tables["user"]} == {"person_1", "person_2"}


def test_excel_export_looks_like_the_old_sheet(client: TestClient, imported: ImportResult) -> None:
    r = client.get("/api/export/budget.xlsx")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.sheetnames == ["Maternity", "Back to normal"]
    ws = wb["Maternity"]
    assert ws.cell(4, 1).value == "Person 1"
    assert ws.cell(4, 5).value == "Person 2"
    assert ws.cell(4, 9).value == "Joint Billing"
    values = [[c.value for c in row] for row in ws.iter_rows()]
    labels = {row[0] for row in values}
    assert {"Transfers", "Surplus", "Household"} <= labels
    household = next(row for row in values if row[0] == "Household")
    assert str(household[2]) == "453.24"
